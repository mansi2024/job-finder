#!/usr/bin/env python3
"""Pick today's jobs to apply to. Usage: python3 apply/queue.py [--limit 5] [--file output/jobs_latest.xlsx]"""
import argparse, csv, json, os, re, sys
import openpyxl
HERE=os.path.dirname(os.path.abspath(__file__)); ROOT=os.path.dirname(HERE)
cfg=json.load(open(os.path.join(HERE,'apply_config.json')))
ap=argparse.ArgumentParser(); ap.add_argument('--limit',type=int,default=cfg['max_per_run'])
ap.add_argument('--new-first',action='store_true',help='rank jobs marked NEW first')
ap.add_argument('--file',default=os.path.join(ROOT,'output','jobs_latest.xlsx')); a=ap.parse_args()
log=os.path.join(HERE,'applied_log.csv'); done_urls=set(); done_keys=set()
if os.path.exists(log):
    for r in csv.DictReader(open(log,encoding='utf8')):
        done_urls.add(r['url']); done_keys.add((r['company'].lower(),r['title'].lower()))
def exp_range(s):
    s=(s or '').lower()
    m=re.search(r'(\d+(?:\.\d+)?)\s*(?:-|to)\s*(\d+(?:\.\d+)?)',s)
    if m: return float(m.group(1)),float(m.group(2))
    m=re.search(r'(\d+(?:\.\d+)?)\s*\+',s)
    if m: return float(m.group(1)),None
    return None
ef=cfg['experience_filter']; wb=openpyxl.load_workbook(a.file); seen=set(); rows=[]
for ws in wb:
    hdr=[c.value for c in ws[1]]
    if 'Apply' not in hdr or 'Status' not in hdr: continue
    ix={h:i for i,h in enumerate(hdr)}
    for r in range(2,ws.max_row+1):
        g=lambda h: ws.cell(row=r,column=ix[h]+1).value
        cell=ws.cell(row=r,column=ix['Apply']+1); url=cell.hyperlink.target if cell.hyperlink else None
        if not url or url in seen: continue
        seen.add(url); title=(g('Title') or ''); company=(g('Company') or '')
        if g('Status') not in cfg['statuses_to_apply']: continue
        if g('Category') not in cfg['categories']: continue
        if any(w in company.lower() for w in cfg.get('company_exclude',[])): continue
        if any(w in title.lower() for w in cfg['title_exclude']): continue
        if url in done_urls or (company.lower(),title.lower()) in done_keys: continue
        loc=(g('Location') or '').lower()
        if any(w in (title+' '+company).lower() for w in cfg['location_exclude']+[' ny, us',', us','in office']): continue
        if (g('Source') or '')=='Y Combinator' and not any(w in (title+' '+loc).lower() for w in ['india','remote']): continue
        if loc:
            if any(w in loc for w in cfg['location_exclude']): continue
            if not any(w in loc for w in cfg['location_allow']): continue
        elif not cfg['accept_missing_location']: continue
        er=exp_range(g('Experience'))
        if er is None:
            if not ef['accept_not_stated']: continue
        else:
            lo,hi=er
            if lo>ef['accept_if_min_years_at_most']: continue
            if hi is not None and hi<ef['accept_if_max_years_at_least']: continue
        is_new=(g('New?')=='NEW')
        score=cfg['rank']['tier'].get(g('Company tier'),0)*10+cfg['rank']['category'].get(g('Category'),0)*3+(g('Skill match') or 0)-(g('Days ago') or 0)*0.2+(1000 if (a.new_first and is_new) else 0)
        rows.append(dict(new=is_new,score=round(score,1),sheet=ws.title,row=r,title=title,company=company,tier=g('Company tier'),location=g('Location'),experience=g('Experience'),skill_match=g('Skill match'),source=g('Source'),url=url))
rows.sort(key=lambda x:-x['score'])
print(json.dumps(rows[:a.limit],indent=2,ensure_ascii=False)); print(f"# {len(rows)} eligible, showing {min(a.limit,len(rows))}",file=sys.stderr)
