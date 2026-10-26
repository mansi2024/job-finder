#!/usr/bin/env python3
"""Record an outcome. Usage: python3 apply/mark.py --url URL --status Applied|Skipped|"Needs review" --note "text" """
import argparse, csv, datetime, os, shutil
import openpyxl
HERE=os.path.dirname(os.path.abspath(__file__)); ROOT=os.path.dirname(HERE)
ap=argparse.ArgumentParser(); ap.add_argument('--url',required=True); ap.add_argument('--status',required=True)
ap.add_argument('--note',default=''); ap.add_argument('--file',default=os.path.join(ROOT,'output','jobs_latest.xlsx')); a=ap.parse_args()
bk=os.path.join(ROOT,'output','backup'); os.makedirs(bk,exist_ok=True)
shutil.copy2(a.file,os.path.join(bk,datetime.datetime.now().strftime('jobs_latest_%Y%m%d_%H%M%S.xlsx')))
wb=openpyxl.load_workbook(a.file); title=company=''; hit=0
for ws in wb:
    hdr=[c.value for c in ws[1]]
    if 'Apply' not in hdr: continue
    ix={h:i+1 for i,h in enumerate(hdr)}
    for r in range(2,ws.max_row+1):
        c=ws.cell(row=r,column=ix['Apply']); 
        if c.hyperlink and c.hyperlink.target==a.url:
            ws.cell(row=r,column=ix['Status']).value=a.status
            n=ws.cell(row=r,column=ix['Notes']); n.value=(f"{datetime.date.today()}: {a.note}").strip(': ')
            title=ws.cell(row=r,column=ix['Title']).value or title; company=ws.cell(row=r,column=ix['Company']).value or company; hit+=1
wb.save(a.file)
log=os.path.join(HERE,'applied_log.csv'); new=not os.path.exists(log)
with open(log,'a',newline='',encoding='utf8') as f:
    w=csv.writer(f)
    if new: w.writerow(['date','company','title','url','status','note'])
    w.writerow([datetime.datetime.now().isoformat(timespec='seconds'),company,title,a.url,a.status,a.note])
print(f"updated {hit} row(s) -> {a.status}")
